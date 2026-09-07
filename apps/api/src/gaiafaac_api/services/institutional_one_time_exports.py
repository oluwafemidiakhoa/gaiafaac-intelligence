from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

from gaiafaac_api.services.branded_one_time_exports import (
    build_one_time_excel as _build_branded_excel,
)
from gaiafaac_api.services.branded_one_time_exports import (
    build_one_time_pdf as _build_branded_pdf,
)

_AMBER = "F7C948"
_TEAL = "07594F"


def _published_series_bounds(sheet) -> tuple[int, int] | None:
    header_row = None
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == "Published period":
            header_row = row
            break
    if header_row is None:
        return None

    last_row = header_row
    for row in range(header_row + 1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value in (None, ""):
            break
        last_row = row
    if last_row == header_row:
        return None
    return header_row, last_row


def _rebuild_visible_analytics_charts(body: bytes) -> bytes:
    workbook = load_workbook(io.BytesIO(body))
    if "Fiscal Analytics" not in workbook.sheetnames:
        return body

    sheet = workbook["Fiscal Analytics"]
    bounds = _published_series_bounds(sheet)
    if bounds is None:
        return body
    header_row, last_row = bounds

    # The original charts used hidden helper columns. Desktop Excel commonly
    # suppresses hidden-cell series, which can leave a valid chart object with
    # a blank plot area. Rebuild the charts from the visible governed evidence
    # table so the workbook renders exactly what the user can audit on-screen.
    sheet._charts = []

    categories = Reference(sheet, min_col=1, min_row=header_row + 1, max_row=last_row)

    trend = LineChart()
    trend.title = "Monthly gross vs net allocation"
    trend.style = 13
    trend.height = 7.4
    trend.width = 13.2
    trend.y_axis.title = "NGN"
    trend.x_axis.title = "Published period"
    for column in (2, 4):
        trend.add_data(
            Reference(
                sheet,
                min_col=column,
                max_col=column,
                min_row=header_row,
                max_row=last_row,
            ),
            titles_from_data=True,
        )
    trend.set_categories(categories)
    if trend.series:
        trend.series[0].graphicalProperties.line.solidFill = _AMBER
    if len(trend.series) > 1:
        trend.series[1].graphicalProperties.line.solidFill = _TEAL
    trend.legend.position = "b"
    sheet.add_chart(trend, "J5")

    burden = BarChart()
    burden.type = "col"
    burden.title = "Deduction burden by published period"
    burden.style = 10
    burden.height = 7.4
    burden.width = 13.2
    burden.y_axis.title = "% of gross"
    burden.y_axis.numFmt = "0%"
    burden.x_axis.title = "Published period"
    burden.add_data(
        Reference(
            sheet,
            min_col=5,
            max_col=5,
            min_row=header_row,
            max_row=last_row,
        ),
        titles_from_data=True,
    )
    burden.set_categories(categories)
    if burden.series:
        burden.series[0].graphicalProperties.solidFill = _TEAL
        burden.series[0].graphicalProperties.line.solidFill = _TEAL
    burden.legend = None
    sheet.add_chart(burden, "J21")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_one_time_excel(**kwargs: Any) -> tuple[str, str, bytes]:
    filename, media_type, body = _build_branded_excel(**kwargs)
    if kwargs.get("product_code") == "decision_pack":
        body = _rebuild_visible_analytics_charts(body)
    return filename, media_type, body


def build_one_time_pdf(**kwargs: Any) -> tuple[str, str, bytes]:
    return _build_branded_pdf(**kwargs)
