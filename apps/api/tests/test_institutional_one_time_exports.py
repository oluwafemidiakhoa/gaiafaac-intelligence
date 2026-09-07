from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from gaiafaac_api.services.institutional_one_time_exports import (
    _rebuild_visible_analytics_charts,
    build_one_time_excel,
)

_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def _artifact() -> dict:
    # Synthetic regression data, not a statement about an actual jurisdiction.
    months = []
    for index, (gross, deductions, net) in enumerate(
        [(100_000, 20_000, 80_000), (120_000, 24_000, 96_000), (90_000, 22_500, 67_500)],
        start=1,
    ):
        period = f"2026-0{index}-01"
        months.append(
            {
                "revenue_month": period,
                "reporting_label": f"Month {index}",
                "gross_total": str(gross),
                "total_deductions": str(deductions),
                "net_allocation": str(net),
                "reconciliation_status": "reconciled",
                "proof_id": f"GF1-TEST-{index}",
                "proof_path": f"/fiscal-proof/lagos/{period}",
                "source_organization": "SYNTHETIC TEST FIXTURE",
                "source_sha256": str(index) * 64,
                "human_verified": True,
            }
        )
    return {
        "schema": "gaia-sample-decision-pack-v1",
        "captured_at": "2026-09-07T05:00:00+00:00",
        "request": {"state_slug": "lagos", "year": 2026, "sample": True},
        "decision_packet": {
            "state_name": "Lagos",
            "year": 2026,
            "coverage_label": "Synthetic three-period test",
            "months_published": 3,
            "annual_gross": "310000",
            "annual_deductions": "66500",
            "annual_net": "243500",
            "deduction_burden_pct": 21.45,
            "net_retention_pct": 78.55,
            "momentum": "Insufficient data",
            "momentum_pct": None,
            "volatility": "Moderate",
            "volatility_cv_pct": 14.2,
            "evidence_status": "Verified",
            "igr_records": [],
            "igr_note": "No IGR supplied in this synthetic fixture.",
            "watch_events": [],
            "months": months,
            "disclaimer": "Synthetic regression fixture only.",
        },
        "statement": "SAMPLE",
    }


def _charts(body: bytes) -> list[ET.Element]:
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        return [
            ET.fromstring(archive.read(name))
            for name in sorted(archive.namelist())
            if name.startswith("xl/charts/chart") and name.endswith(".xml")
        ]


@pytest.mark.parametrize("sample", [True, False])
def test_samples_and_paid_exports_have_visible_references_and_populated_caches(sample):
    _, _, body = build_one_time_excel(
        purchase_id="SAMPLE-lagos-2026" if sample else "12345678-1234-5678-1234-567812345678",
        product_code="decision_pack",
        amount_naira="50000",
        currency="NGN",
        completed_at="2026-09-07T05:00:00+00:00",
        artifact=_artifact(),
        sample=sample,
        jurisdiction="Lagos",
    )
    charts = _charts(body)
    assert len(charts) == 2
    formulas = [node.text for chart in charts for node in chart.findall(".//c:f", _NS)]
    for column in ("A", "B", "D", "E"):
        assert any(f"'Fiscal Analytics'!${column}$" in formula for formula in formulas)
    for column in ("X", "Y", "Z", "AA"):
        assert all(f"!${column}$" not in formula for formula in formulas)
    for chart in charts:
        assert chart.find(".//c:plotVisOnly", _NS).get("val") == "0"
        assert chart.find(".//c:dispBlanksAs", _NS).get("val") == "gap"
        for series in chart.findall(".//c:ser", _NS):
            assert len(series.findall("c:cat/c:strRef/c:strCache/c:pt", _NS)) == 3
            assert len(series.findall("c:val/c:numRef/c:numCache/c:pt", _NS)) == 3
    workbook = load_workbook(io.BytesIO(body))
    assert workbook["Document Control"]["A1"].value == "GAIA FISCAL INTELLIGENCE"


def test_chart_cache_preserves_zero_but_does_not_turn_missing_values_into_zero():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fiscal Analytics"
    sheet.append(["Published period", "Gross", "Deductions", "Net", "Burden"])
    sheet.append(["Jan 2026", 100, 20, 80, 0.2])
    sheet.append(["Feb 2026", 100, None, None, None])
    sheet.append(["Mar 2026", 0, 0, 0, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    charts = _charts(_rebuild_visible_analytics_charts(buffer.getvalue()))
    net_cache = charts[0].findall(".//c:ser", _NS)[1].find("c:val/c:numRef/c:numCache", _NS)
    points = net_cache.findall("c:pt", _NS)
    assert [point.get("idx") for point in points] == ["0", "2"]
    assert float(points[1].find("c:v", _NS).text) == 0
    assert net_cache.find("c:ptCount", _NS).get("val") == "3"
